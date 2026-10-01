"""Google Drive, OneDrive and Outlook Calendar: transport, matching, routes.

Nothing here reaches the network. ``requests.request`` / ``.get`` / ``.post``
are replaced with a router that answers by URL and records every call, so
the tests can assert on the exact request each provider would have received
(the query string, the multipart body, the Authorization header) without
either provider existing.

The unconfigured-route tests matter as much as the configured ones: the
owner had not created the OAuth apps when this shipped, so every route has to
answer "not set up yet" cleanly rather than 500.
"""

from __future__ import annotations

import io
import json
import time
import zipfile
from datetime import date, datetime, timezone
from urllib.parse import parse_qs, urlparse

import pytest
import requests

import App
import cloud_token_client
import google_drive_helper as gd
import onedrive_helper as od
import outlook_calendar_helper as outlook
from App import CloudDocsIntegration, OutlookIntegration, User, db
from intelliplan.services import document_matcher as dm


# ── Fakes ─────────────────────────────────────────────────────────────


class FakeResponse:
    def __init__(self, payload=None, status=200, body=None, headers=None):
        self._payload = payload
        self.status_code = status
        self._body = body if body is not None else (
            json.dumps(payload).encode() if payload is not None else b"")
        self.headers = headers or {}
        self.text = self._body.decode("utf-8", errors="replace")

    def json(self):
        if self._payload is None:
            raise ValueError("no json")
        return self._payload

    def raise_for_status(self):
        if self.status_code >= 400:
            raise requests.HTTPError(f"HTTP {self.status_code}", response=self)

    def iter_content(self, chunk_size=65536):
        for i in range(0, len(self._body), chunk_size):
            yield self._body[i:i + chunk_size]

    def __enter__(self):
        return self

    def __exit__(self, *exc):
        return False


class Router:
    """Answers HTTP calls by (method, url-substring); records all of them."""

    def __init__(self):
        self.routes = []
        self.calls = []

    def add(self, method, fragment, response):
        self.routes.append((method.upper(), fragment, response))

    def __call__(self, method, url, **kwargs):
        self.calls.append({"method": method.upper(), "url": url, **kwargs})
        for m, fragment, response in self.routes:
            if m == method.upper() and fragment in url:
                return response(url, kwargs) if callable(response) else response
        raise AssertionError(f"unexpected {method} {url}")

    def install(self, monkeypatch):
        monkeypatch.setattr(requests, "request", self)
        monkeypatch.setattr(requests, "get", lambda url, **kw: self("GET", url, **kw))
        monkeypatch.setattr(requests, "post", lambda url, **kw: self("POST", url, **kw))
        return self

    def find(self, method, fragment):
        return [c for c in self.calls if c["method"] == method and fragment in c["url"]]


@pytest.fixture
def router(monkeypatch):
    return Router().install(monkeypatch)


def _docx_bytes(*paragraphs):
    from docx import Document
    doc = Document()
    for p in paragraphs:
        doc.add_paragraph(p)
    out = io.BytesIO()
    doc.save(out)
    return out.getvalue()


def _pptx_bytes(*slides):
    out = io.BytesIO()
    with zipfile.ZipFile(out, "w") as z:
        for i, text in enumerate(slides, 1):
            z.writestr(f"ppt/slides/slide{i}.xml",
                       f"<p:sld><a:t>{text}</a:t><a:t>&amp; more</a:t></p:sld>")
    return out.getvalue()


@pytest.fixture(autouse=True)
def _fresh_cache():
    dm.clear_cache()
    yield
    dm.clear_cache()


# ── Search query building ─────────────────────────────────────────────


def test_drive_query_ors_terms_filters_to_readable_types_and_excludes_trash():
    q = gd.build_search_query(["photosynthesis", "chlorophyll"])
    assert "fullText contains 'photosynthesis' or fullText contains 'chlorophyll'" in q
    assert "mimeType = 'application/vnd.google-apps.document'" in q
    assert "mimeType = 'application/pdf'" in q
    assert q.endswith("trashed = false")


def test_drive_query_escapes_quotes_and_backslashes():
    """An apostrophe in a title would otherwise end the literal and turn the
    whole search into a 400."""
    q = gd.build_search_query(["Frankenstein's creature", "a\\b"])
    assert "fullText contains 'Frankenstein\\'s creature'" in q
    assert "fullText contains 'a\\\\b'" in q


def test_drive_query_rejects_empty_and_dedupes():
    with pytest.raises(ValueError):
        gd.build_search_query(["  ", ""])
    q = gd.build_search_query(["Cell", "cell", "mitosis"])
    assert q.count("fullText contains") == 2


def test_onedrive_search_path_doubles_quotes_and_percent_encodes():
    path = od.search_path("Romeo & Juliet's balcony")
    assert path.startswith("/me/drive/root/search(q='")
    assert "Juliet%27%27s" in path  # '' then encoded
    assert "%26" in path and " " not in path


def test_drive_search_normalises_results(router):
    router.add("GET", "/drive/v3/files", FakeResponse({"files": [
        {"id": "doc_12345", "name": "Bio notes", "mimeType": gd.GOOGLE_DOC,
         "modifiedTime": "2026-09-01T10:00:00.000Z", "webViewLink": "https://docs.google.com/d/1"},
        {"name": "no id, dropped"},
    ]}))
    client = gd.DriveClient({"token": "t", "expires_at": time.time() + 3600})
    files = gd.search_files(client, ["biology"])
    assert files == [{"provider": "google_drive", "id": "doc_12345", "name": "Bio notes",
                      "mime": gd.GOOGLE_DOC, "modified": "2026-09-01T10:00:00.000Z",
                      "url": "https://docs.google.com/d/1", "size": 0}]
    call = router.calls[0]
    assert call["params"]["q"].startswith("(fullText contains 'biology')")
    assert "orderBy" not in call["params"]  # Drive refuses to sort fullText queries
    assert call["headers"]["Authorization"] == "Bearer t"


# ── Text extraction ───────────────────────────────────────────────────


def test_google_doc_is_exported_as_plain_text(router):
    router.add("GET", "/files/doc_12345/export",
               FakeResponse(body="Photosynthesis   turns light\ninto sugar".encode()))
    client = gd.DriveClient({"token": "t", "expires_at": time.time() + 3600})
    text = gd.fetch_text(client, {"id": "doc_12345", "mime": gd.GOOGLE_DOC, "name": "Notes"})
    assert text == "Photosynthesis turns light into sugar"
    assert router.calls[0]["params"] == {"mimeType": "text/plain"}


def test_drive_docx_is_downloaded_and_parsed(router):
    router.add("GET", "/files/file_12345", FakeResponse(body=_docx_bytes("Mitosis has four phases.")))
    client = gd.DriveClient({"token": "t", "expires_at": time.time() + 3600})
    item = {"id": "file_12345", "name": "cells.docx",
            "mime": "application/vnd.openxmlformats-officedocument.wordprocessingml.document", "size": 5000}
    assert gd.fetch_text(client, item) == "Mitosis has four phases."
    assert router.calls[0]["params"] == {"alt": "media"}


def test_unreadable_and_oversized_files_are_skipped_without_download(router):
    client = gd.DriveClient({"token": "t", "expires_at": time.time() + 3600})
    assert gd.fetch_text(client, {"id": "img_12345", "name": "photo.jpg", "mime": "image/jpeg"}) == ""
    assert gd.fetch_text(client, {"id": "big_12345", "name": "book.pdf", "mime": "application/pdf",
                                  "size": gd.MAX_DOWNLOAD_BYTES + 1}) == ""
    assert router.calls == []


def test_streamed_body_over_the_cap_is_abandoned(router):
    router.add("GET", "/files/pdf_12345", FakeResponse(body=b"x" * (gd.MAX_DOWNLOAD_BYTES + 10)))
    client = gd.DriveClient({"token": "t", "expires_at": time.time() + 3600})
    assert gd.fetch_text(client, {"id": "pdf_12345", "name": "a.txt", "mime": "text/plain"}) == ""


def test_onedrive_pptx_text_is_extracted(router):
    router.add("GET", "/me/drive/items/ABC123/content", FakeResponse(body=_pptx_bytes("Cell cycle", "Mitosis")))
    client = od.OneDriveClient({"access_token": "t", "expires_at": time.time() + 3600})
    text = od.fetch_text(client, {"id": "ABC123", "name": "Unit 4.pptx", "mime": "", "size": 900})
    assert text == "Cell cycle & more Mitosis & more"


def test_onedrive_shared_item_is_read_from_its_own_drive(router):
    router.add("GET", "/drives/drv1/items/ITEM9/content", FakeResponse(body=b"shared notes"))
    client = od.OneDriveClient({"access_token": "t", "expires_at": time.time() + 3600})
    assert od.fetch_text(client, {"id": "ITEM9", "drive_id": "drv1", "name": "n.txt"}) == "shared notes"


def test_item_ids_cannot_escape_the_api_path():
    client = od.OneDriveClient({"access_token": "t"})
    with pytest.raises(cloud_token_client.ProviderError):
        od.fetch_text(client, {"id": "../../me/messages", "name": "x.txt"})
    with pytest.raises(cloud_token_client.ProviderError):
        gd.fetch_text(gd.DriveClient({"token": "t"}), {"id": "a/b?c", "name": "x.txt"})


# ── Matcher ───────────────────────────────────────────────────────────


def test_keywords_drop_stopwords_and_classroom_boilerplate_title_first():
    kws = dm.extract_keywords("Photosynthesis Lab Report assignment", "AP Biology - Period 3",
                              "Due Friday. Explain photosynthesis and chlorophyll. chlorophyll absorbs light.")
    assert kws[0] == "photosynthesis"
    lowered = [k.lower() for k in kws]
    for noise in ("assignment", "period", "due", "the", "3"):
        assert noise not in lowered
    assert "biology" in lowered
    assert "chlorophyll" in lowered  # repeated in the directions


def test_name_hits_outrank_recency():
    now = datetime(2026, 9, 30, tzinfo=timezone.utc)
    ranked = dm.rank([
        {"provider": "google_drive", "id": "new", "name": "Random doc", "modified": "2026-09-29T00:00:00Z"},
        {"provider": "google_drive", "id": "old", "name": "Photosynthesis notes", "modified": "2025-01-01T00:00:00Z"},
    ], ["photosynthesis"], now=now)
    assert [r["id"] for r in ranked] == ["old", "new"]


def test_recency_breaks_ties_and_duplicates_collapse():
    now = datetime(2026, 9, 30, tzinfo=timezone.utc)
    ranked = dm.rank([
        {"provider": "onedrive", "id": "a", "name": "Essay draft", "modified": "2026-01-01T00:00:00Z"},
        {"provider": "onedrive", "id": "b", "name": "Essay outline", "modified": "2026-09-28T00:00:00.1234567Z"},
        {"provider": "onedrive", "id": "b", "name": "Essay outline", "modified": "2026-09-28T00:00:00Z"},
    ], ["essay"], now=now)
    assert [r["id"] for r in ranked] == ["b", "a"]


def _source(provider, files, texts, calls=None):
    def search(keywords, course):
        if calls is not None:
            calls.append(list(keywords))
        return [dict(f) for f in files]
    return dm.Source(provider, search, lambda item: texts.get(item["id"], ""))


def test_find_documents_reads_top_matches_and_drops_unrelated_text():
    now = datetime(2026, 9, 30, tzinfo=timezone.utc)
    drive = _source("google_drive", [
        {"id": "1", "name": "Photosynthesis notes", "modified": "2026-09-20T00:00:00Z", "url": "https://d/1"},
        {"id": "2", "name": "Scan 0042", "modified": "2026-09-25T00:00:00Z"},
    ], {"1": "Chlorophyll absorbs light during photosynthesis.", "2": "Grocery list: eggs, milk"})
    result = dm.find_documents([drive], "Photosynthesis lab", "Biology", now=now)
    assert [d["id"] for d in result["documents"]] == ["1"]
    doc = result["documents"][0]
    assert doc["provider_label"] == "Google Drive"
    assert "photosynthesis" in doc["excerpt"].lower()
    assert "text" not in dm.public_documents(result["documents"])[0]


def test_one_failing_provider_does_not_hide_the_other():
    def broken(keywords, course):
        raise cloud_token_client.TokenExpired("gone")
    bad = dm.Source("onedrive", broken, lambda item: "")
    good = _source("google_drive", [{"id": "1", "name": "Essay plan"}], {"1": "essay plan text"})
    result = dm.find_documents([bad, good], "Essay plan", "")
    assert [d["id"] for d in result["documents"]] == ["1"]
    assert result["errors"] == [{"provider": "onedrive", "reason": "reconnect"}]


def test_results_are_cached_per_owner_so_chat_turns_do_not_research():
    calls = []
    src = _source("google_drive", [{"id": "1", "name": "Essay plan"}], {"1": "essay"}, calls)
    dm.find_documents([src], "Essay plan", owner=7)
    dm.find_documents([src], "Essay plan", owner=7)
    assert len(calls) == 1
    dm.find_documents([src], "Essay plan", owner=8)
    assert len(calls) == 2


def test_augment_context_keeps_canvas_first_and_labels_student_files():
    ctx = {"title": "Lab", "description": "", "materials": [{"name": "rubric.pdf", "text": "rubric"}],
           "skipped_count": 0}
    out = dm.augment_context(ctx, [{"provider": "onedrive", "provider_label": "OneDrive",
                                    "name": "notes.docx", "text": "x" * 20000}], max_total_chars=500)
    assert out["materials"][0]["name"] == "rubric.pdf"
    assert out["materials"][1]["name"] == "Your OneDrive: notes.docx"
    assert len(out["materials"][1]["text"]) == 500
    from assignment_materials import assignment_prompt
    assert "student's own files" in assignment_prompt(out)


def test_study_guide_spreads_steps_to_the_due_date_and_links_documents():
    guide = dm.build_study_guide(
        "Photosynthesis lab", "Biology", "Explain the light reactions.", "2026-10-10",
        documents=[{"provider": "google_drive", "provider_label": "Google Drive", "name": "Notes",
                    "url": "https://docs.google.com/d/1", "modified": "2026-09-01T00:00:00Z",
                    "text": "photosynthesis chlorophyll chlorophyll thylakoid thylakoid"}],
        today=date(2026, 9, 30))
    md = guide["markdown"]
    assert md.startswith("# Study guide: Photosynthesis lab")
    assert "- [ ] Wed Sep 30:" in md and "- [ ] Fri Oct 09:" in md
    assert "[Notes](https://docs.google.com/d/1)" in md
    assert "**chlorophyll**" in md
    assert guide["title"] == "Photosynthesis lab study guide"


# ── Create-doc request shape ──────────────────────────────────────────


def test_drive_create_makes_folder_once_then_uploads_a_converted_google_doc(router):
    # Upload first: its URL also contains "/drive/v3/files".
    router.add("POST", "/upload/drive/v3/files", FakeResponse(
        {"id": "doc_9", "name": "Lab study guide", "webViewLink": "https://docs.google.com/d/9"}))
    router.add("GET", "/drive/v3/files", FakeResponse({"files": []}))
    router.add("POST", "/drive/v3/files", FakeResponse({"id": "folder_1"}))
    client = gd.DriveClient({"token": "t", "expires_at": time.time() + 3600})
    created = gd.create_document(client, "Lab study guide", "# Lab\n- [ ] Step **one**")
    assert created == {"provider": "google_drive", "id": "doc_9", "name": "Lab study guide",
                       "url": "https://docs.google.com/d/9"}

    folder_query = router.find("GET", "/drive/v3/files")[0]["params"]["q"]
    assert "appProperties has { key='intelliplan' and value='folder' }" in folder_query
    folder_create = [c for c in router.find("POST", "/drive/v3/files") if "upload" not in c["url"]][0]["json"]
    assert folder_create["mimeType"] == gd.FOLDER and folder_create["appProperties"] == {"intelliplan": "folder"}

    upload = router.find("POST", "/upload/drive/v3/files")[0]
    assert upload["params"]["uploadType"] == "multipart"
    ctype = upload["headers"]["Content-Type"]
    assert ctype.startswith("multipart/related; boundary=")
    body = upload["data"].decode()
    metadata = json.loads(body.split("\r\n\r\n", 1)[1].split("\r\n", 1)[0])
    assert metadata == {"name": "Lab study guide", "mimeType": gd.GOOGLE_DOC, "parents": ["folder_1"]}
    assert "<h1>Lab</h1>" in body and "<b>one</b>" in body


def test_onedrive_create_puts_a_docx_into_the_intelliplan_folder(router):
    router.add("PUT", "/me/drive/root:/IntelliPlan/", FakeResponse(
        {"id": "od1", "name": "Essay study guide.docx", "webUrl": "https://onedrive.live.com/x"}))
    client = od.OneDriveClient({"access_token": "t", "expires_at": time.time() + 3600})
    created = od.create_document(client, 'Essay: "Hamlet"? study guide', "# Hamlet\n- point")
    assert created["url"] == "https://onedrive.live.com/x"
    call = router.calls[0]
    assert call["url"].endswith(":/content?@microsoft.graph.conflictBehavior=rename")
    assert "Essay%20Hamlet%20study%20guide.docx" in call["url"]  # unsafe chars stripped
    assert call["headers"]["Content-Type"].endswith("wordprocessingml.document")
    assert zipfile.ZipFile(io.BytesIO(call["data"])).namelist()


# ── Token refresh ─────────────────────────────────────────────────────


def test_expired_drive_token_is_refreshed_before_the_call_and_persisted(router, monkeypatch):
    monkeypatch.setenv("GOOGLE_CLIENT_ID", "cid")
    monkeypatch.setenv("GOOGLE_CLIENT_SECRET", "csecret")
    router.add("POST", "oauth2.googleapis.com/token", FakeResponse({"access_token": "fresh", "expires_in": 3600}))
    router.add("GET", "/drive/v3/files", FakeResponse({"files": []}))
    saved = []
    client = gd.DriveClient({"token": "stale", "refresh_token": "r1", "expires_at": time.time() - 5},
                            on_refresh=saved.append)
    gd.search_files(client, ["x-ray"])
    token_call = router.find("POST", "oauth2.googleapis.com/token")[0]
    assert token_call["data"]["grant_type"] == "refresh_token"
    assert token_call["data"]["refresh_token"] == "r1"
    assert router.find("GET", "/drive/v3/files")[0]["headers"]["Authorization"] == "Bearer fresh"
    assert saved and saved[0]["token"] == "fresh" and saved[0]["refresh_token"] == "r1"
    assert saved[0]["expires_at"] > time.time() + 3000


def test_a_401_triggers_one_refresh_and_retry(router, monkeypatch):
    monkeypatch.setenv("MICROSOFT_CLIENT_ID", "mid")
    monkeypatch.setenv("MICROSOFT_CLIENT_SECRET", "msecret")
    answers = iter([FakeResponse({}, status=401), FakeResponse({"value": []})])
    router.add("GET", "/search(q=", lambda url, kw: next(answers))
    router.add("POST", "login.microsoftonline.com", FakeResponse(
        {"access_token": "new", "refresh_token": "rotated", "expires_in": 3600}))
    client = od.OneDriveClient({"access_token": "old", "refresh_token": "r0"})
    assert od.search_files(client, "essay") == []
    token_call = router.find("POST", "login.microsoftonline.com")[0]
    # The refresh asks for the Files scopes this token was granted, not the
    # calendar ones: asking for unconsented scopes is an invalid_grant.
    assert "Files.ReadWrite" in token_call["data"]["scope"]
    assert "Calendars.ReadWrite" not in token_call["data"]["scope"]
    assert client.token["refresh_token"] == "rotated"
    assert router.find("GET", "/search(q=")[1]["headers"]["Authorization"] == "Bearer new"


def test_a_rejected_refresh_raises_token_expired_without_the_token(router, monkeypatch):
    router.add("POST", "oauth2.googleapis.com/token", FakeResponse(
        {"error": "invalid_grant", "echo": "SECRET-REFRESH"}, status=400))
    client = gd.DriveClient({"token": "t", "refresh_token": "SECRET-REFRESH", "expires_at": 1})
    with pytest.raises(cloud_token_client.TokenExpired) as err:
        client.request("GET", gd.API + "/files")
    assert "SECRET-REFRESH" not in str(err.value)


# ── Outlook Calendar ──────────────────────────────────────────────────


def test_busy_window_is_the_students_local_days_in_utc_and_free_events_are_ignored(monkeypatch):
    seen = {}

    def fake_get(token, path, params=None):
        seen.update(params or {})
        return {"value": [
            {"start": {"dateTime": "2026-03-02T18:00:00.0000000"}, "end": {"dateTime": "2026-03-02T19:00:00.0000000"},
             "showAs": "busy"},
            {"start": {"dateTime": "2026-03-02T20:00:00"}, "end": {"dateTime": "2026-03-02T21:00:00"},
             "showAs": "free"},
            {"start": {"dateTime": "2026-03-02T21:00:00"}, "end": {"dateTime": "2026-03-02T22:00:00"},
             "isCancelled": True},
        ]}
    monkeypatch.setattr(outlook, "graph_get", fake_get)
    busy = outlook.busy_minutes_by_date({"access_token": "x"}, date(2026, 3, 2), days=1,
                                        utc_offset_minutes=-7 * 60)
    # Local midnight at UTC-7 is 07:00 UTC.
    assert seen["startDateTime"] == "2026-03-02T07:00:00Z"
    assert seen["endDateTime"] == "2026-03-03T07:00:00Z"
    assert busy == {date(2026, 3, 2): [(11 * 60, 12 * 60)]}


def test_calendar_view_follows_paging(monkeypatch):
    pages = {
        None: {"value": [{"id": "1"}], "@odata.nextLink": outlook.GRAPH + "/me/calendarView?$skip=1"},
        outlook.GRAPH + "/me/calendarView?$skip=1": {"value": [{"id": "2"}]},
    }
    monkeypatch.setattr(outlook, "graph_get",
                        lambda token, path, params=None: pages[None if path == "/me/calendarView" else path])
    events = outlook._calendar_view({"access_token": "x"}, datetime(2026, 1, 1), datetime(2026, 1, 2))
    assert [e["id"] for e in events] == ["1", "2"]


def test_export_is_idempotent_and_uses_the_students_zone(monkeypatch):
    existing = {"value": [{"subject": "Study: Essay", "start": {"dateTime": "2026-03-02T23:30:00"},
                           "end": {"dateTime": "2026-03-03T00:30:00"}, "showAs": "busy"}]}
    monkeypatch.setattr(outlook, "graph_get", lambda *a, **k: existing)
    posted = []
    monkeypatch.setattr(outlook, "graph_post", lambda token, path, body: posted.append(body) or {"id": "new"})
    schedule = {"schedule": [{"date": "2026-03-02", "blocks": [
        {"assignment": "Essay", "course": "English", "time_slot": "4:30 PM - 5:30 PM", "duration_minutes": 60},
        {"assignment": "Math", "course": "Algebra", "time_slot": "18:00 - 18:30", "duration_minutes": 30},
    ]}]}
    result = outlook.export_schedule({"access_token": "x"}, schedule, tz_name="America/Los_Angeles")
    # 4:30 PM PST is 00:30 UTC... the existing Essay block at 23:30 UTC is a
    # different start, so check the Math block's conversion and the duplicate
    # rule separately.
    assert result["created"] == ["new", "new"]
    math = posted[1]
    assert math["start"] == {"dateTime": "2026-03-03T02:00:00", "timeZone": "UTC"}
    assert math["categories"] == ["IntelliPlan"]

    posted.clear()
    existing["value"] = [{"subject": "Study: Math", "start": {"dateTime": "2026-03-03T02:00:00"},
                          "end": {"dateTime": "2026-03-03T02:30:00"}}]
    result = outlook.export_schedule({"access_token": "x"}, schedule, tz_name="America/Los_Angeles")
    assert result["skipped"] == 1 and len(posted) == 1


def test_export_can_skip_blocks_that_overlap_real_events(monkeypatch):
    monkeypatch.setattr(outlook, "graph_get", lambda *a, **k: {"value": [
        {"subject": "Dentist", "start": {"dateTime": "2026-03-02T16:45:00"},
         "end": {"dateTime": "2026-03-02T17:15:00"}, "showAs": "busy"}]})
    monkeypatch.setattr(outlook, "graph_post", lambda *a: pytest.fail("should have been skipped"))
    result = outlook.export_schedule({"access_token": "x"}, {"schedule": [{"date": "2026-03-02", "blocks": [
        {"assignment": "Essay", "time_slot": "4:30 PM - 5:30 PM", "duration_minutes": 60}]}]},
        utc_offset_minutes=0, skip_overlaps=True)
    assert result == {"created": [], "skipped": 1}


def test_upcoming_events_match_the_google_shape(monkeypatch):
    monkeypatch.setattr(outlook, "graph_get", lambda *a, **k: {"value": [
        {"id": "e1", "subject": "Practice", "start": {"dateTime": "2026-10-01T15:00:00.0000000"},
         "end": {"dateTime": "2026-10-01T16:00:00.0000000"}}]})
    assert outlook.get_upcoming_events({"access_token": "x"}) == [{
        "id": "e1", "title": "Practice", "start": "2026-10-01T15:00:00Z", "end": "2026-10-01T16:00:00Z",
        "description": "", "source": "outlook_calendar"}]


def test_redirect_uri_defaults_to_the_app_route(monkeypatch):
    monkeypatch.delenv("MICROSOFT_REDIRECT_URI", raising=False)
    monkeypatch.delenv("APP_BASE_URL", raising=False)
    assert outlook.redirect_uri() == "https://intelliplan.tech/oauth/outlook/callback"
    monkeypatch.setenv("MICROSOFT_CLIENT_ID", "mid")
    monkeypatch.setenv("MICROSOFT_CLIENT_SECRET", "ms")
    assert outlook.configured()


# ── Routes ────────────────────────────────────────────────────────────

CLOUD_ENV = ("GOOGLE_DRIVE_ENABLED", "ONEDRIVE_ENABLED", "GOOGLE_DRIVE_SCOPE_MODE",
             "MICROSOFT_CLIENT_ID", "MICROSOFT_CLIENT_SECRET", "MICROSOFT_REDIRECT_URI",
             "GOOGLE_OAUTH_UNVERIFIED")


@pytest.fixture
def unconfigured(monkeypatch):
    for name in CLOUD_ENV:
        monkeypatch.delenv(name, raising=False)


@pytest.fixture
def configured(monkeypatch):
    monkeypatch.setenv("GOOGLE_CLIENT_ID", "gid")
    monkeypatch.setenv("GOOGLE_CLIENT_SECRET", "gsecret")
    monkeypatch.setenv("GOOGLE_DRIVE_ENABLED", "1")
    monkeypatch.setenv("MICROSOFT_CLIENT_ID", "mid")
    monkeypatch.setenv("MICROSOFT_CLIENT_SECRET", "msecret")
    monkeypatch.setenv("ONEDRIVE_ENABLED", "1")
    monkeypatch.delenv("GOOGLE_OAUTH_UNVERIFIED", raising=False)
    monkeypatch.delenv("GOOGLE_DRIVE_SCOPE_MODE", raising=False)


def _wipe():
    ids = [u.id for u in User.query.filter(User.email.like("clouddocs+%")).all()]
    if ids:
        CloudDocsIntegration.query.filter(CloudDocsIntegration.user_id.in_(ids)).delete(synchronize_session=False)
        OutlookIntegration.query.filter(OutlookIntegration.user_id.in_(ids)).delete(synchronize_session=False)
        App.LinkedAccount.query.filter(App.LinkedAccount.user_id.in_(ids)).delete(synchronize_session=False)
    User.query.filter(User.email.like("clouddocs+%")).delete(synchronize_session=False)
    db.session.commit()


@pytest.fixture
def client():
    App.app.config["TESTING"] = True
    App.limiter.enabled = False
    with App.app.test_client() as c:
        with App.app.app_context():
            db.create_all()
            _wipe()
        yield c
        with App.app.app_context():
            _wipe()
    App.limiter.enabled = True


@pytest.fixture
def uid(client):
    with App.app.app_context():
        u = User(email="clouddocs+a@example.com", name="Cloud Tester",
                 password_hash=App.bcrypt.generate_password_hash("hunter2ok").decode())
        u.ai_personalization_opt_in = True
        u.birth_year = datetime.utcnow().year - 16
        db.session.add(u)
        db.session.commit()
        App._get_or_create_identity(u.id).completed = True
        db.session.commit()
        user_id = u.id
    with client.session_transaction() as s:
        s["_user_id"] = str(user_id)
        s["_fresh"] = True
    return user_id


def _connect(user_id, provider, token):
    with App.app.app_context():
        db.session.add(CloudDocsIntegration(user_id=user_id, provider=provider,
                                            token_data=json.dumps(token), account_email="s@example.com"))
        db.session.commit()


def _stored_token(user_id, provider):
    with App.app.app_context():
        row = CloudDocsIntegration.query.filter_by(user_id=user_id, provider=provider).first()
        return json.loads(row.token_data) if row else None


def test_token_column_is_encrypted_at_rest():
    import secret_box
    assert isinstance(CloudDocsIntegration.__table__.c.token_data.type, secret_box.EncryptedText)


# Unconfigured: every route answers, none 500s.

def test_status_reports_unconfigured_providers(client, unconfigured, uid):
    data = client.get("/api/cloud-docs/status").get_json()
    assert {p["id"]: p["configured"] for p in data["providers"]} == {"google_drive": False, "onedrive": False}


def test_connect_links_redirect_back_to_settings_when_unconfigured(client, unconfigured, uid):
    for path, outcome in (("/oauth/google-drive", "google_drive_unavailable"),
                          ("/oauth/onedrive", "onedrive_unavailable"),
                          ("/oauth/outlook", "outlook_unavailable")):
        r = client.get(path)
        assert r.status_code == 302, path
        assert outcome in r.headers["Location"]


def test_match_and_create_degrade_when_unconfigured(client, unconfigured, uid):
    r = client.get("/api/cloud-docs/match?title=Essay")
    assert r.status_code == 200
    assert r.get_json()["documents"] == [] and r.get_json()["connected"] == []
    r = client.post("/api/cloud-docs/create", json={"title": "Essay"})
    assert r.status_code == 409 and r.get_json()["connect_required"]
    r = client.post("/api/cloud-docs/create", json={"title": "Essay", "provider": "onedrive"})
    assert r.status_code == 503


def test_an_existing_row_is_ignored_once_the_owner_switches_a_provider_off(client, unconfigured, uid, router):
    _connect(uid, "google_drive", {"token": "t", "refresh_token": "r"})
    r = client.get("/api/cloud-docs/match?title=Essay")
    assert r.status_code == 200 and r.get_json()["connected"] == []
    assert router.calls == []


def test_catalog_shows_unconfigured_providers_as_coming_soon(client, unconfigured):
    items = {i["id"]: i for i in client.get("/api/integrations/status").get_json()["integrations"]}
    for key in ("google_drive", "onedrive", "outlook_calendar"):
        assert items[key]["coming_soon"] is True and items[key]["methods"] == []


def test_settings_and_dashboard_render(client, unconfigured, uid):
    page = client.get("/settings")
    assert page.status_code == 200
    html = page.get_data(as_text=True)
    assert 'id="gdriveStatus"' in html and 'id="onedriveStatus"' in html and 'id="outlookStatus"' in html
    dash = client.get("/dashboard")
    assert dash.status_code == 200
    assert 'id="adDocs"' in dash.get_data(as_text=True)


def test_routes_require_sign_in(client, configured):
    assert client.get("/api/cloud-docs/match?title=x").status_code == 401
    assert client.post("/api/cloud-docs/create", json={"title": "x"}).status_code == 401
    assert client.post("/api/cloud-docs/disconnect/google_drive").status_code == 401
    assert client.get("/oauth/google-drive").status_code == 302


def test_bad_input_is_a_400(client, configured, uid):
    assert client.get("/api/cloud-docs/match").status_code == 400
    assert client.post("/api/cloud-docs/create", json={"title": "x", "provider": "dropbox"}).status_code == 400
    assert client.post("/api/cloud-docs/disconnect/dropbox").status_code == 400


# Configured, with mocked provider APIs.

def test_google_drive_connect_asks_for_drive_scopes_incrementally(client, configured, uid):
    r = client.get("/oauth/google-drive")
    assert r.status_code == 302
    query = parse_qs(urlparse(r.headers["Location"]).query)
    scopes = query["scope"][0].split()
    assert gd.SCOPE_FILE in scopes and gd.SCOPE_READONLY in scopes
    assert "https://www.googleapis.com/auth/documents" not in scopes
    assert query["include_granted_scopes"] == ["true"]
    assert query["client_id"] == ["gid"]
    with client.session_transaction() as s:
        assert s["oauth_purpose"] == "drive"


def test_file_scope_mode_drops_the_restricted_scope(client, configured, uid, monkeypatch):
    monkeypatch.setenv("GOOGLE_DRIVE_SCOPE_MODE", "file")
    scopes = parse_qs(urlparse(client.get("/oauth/google-drive").headers["Location"]).query)["scope"][0]
    assert gd.SCOPE_FILE in scopes and gd.SCOPE_READONLY not in scopes


def _google_callback(client, monkeypatch, scopes):
    with client.session_transaction() as s:
        s["oauth_state"] = "st"
        s["oauth_purpose"] = "drive"
        s["oauth_code_verifier"] = "v"
    monkeypatch.setattr(App, "exchange_code_for_token", lambda code, code_verifier=None: {
        "token": "gtok", "refresh_token": "grefresh", "expires_at": time.time() + 3600, "scopes": scopes})
    monkeypatch.setattr(requests, "get", lambda url, **kw: FakeResponse(
        {"sub": "g-1", "email": "student@gmail.com", "name": "Student"}))
    return client.get("/oauth2callback?state=st&code=c")


def test_google_callback_stores_a_drive_grant_in_its_own_row(client, configured, uid, monkeypatch):
    r = _google_callback(client, monkeypatch, [gd.SCOPE_FILE, gd.SCOPE_READONLY, "openid"])
    assert r.status_code == 302 and "google_drive_connected" in r.headers["Location"]
    token = _stored_token(uid, "google_drive")
    assert token["refresh_token"] == "grefresh"
    with App.app.app_context():
        # A Drive grant must not masquerade as a calendar connection.
        assert App.GoogleIntegration.query.filter_by(user_id=uid).count() == 0


def test_google_callback_without_the_drive_box_ticked_stores_nothing(client, configured, uid, monkeypatch):
    r = _google_callback(client, monkeypatch, ["openid", "email"])
    assert "google_drive_denied" in r.headers["Location"]
    assert _stored_token(uid, "google_drive") is None


def test_declining_on_googles_screen_returns_to_settings(client, configured, uid):
    with client.session_transaction() as s:
        s["oauth_state"] = "st"
        s["oauth_purpose"] = "drive"
    r = client.get("/oauth2callback?error=access_denied&state=st")
    assert "google_drive_cancelled" in r.headers["Location"]


def test_onedrive_connect_and_callback_share_the_outlook_redirect_uri(client, configured, uid, router):
    r = client.get("/oauth/onedrive")
    location = urlparse(r.headers["Location"])
    query = parse_qs(location.query)
    assert location.netloc == "login.microsoftonline.com"
    assert "Files.ReadWrite" in query["scope"][0] and "offline_access" in query["scope"][0]
    assert query["redirect_uri"] == ["https://intelliplan.tech/oauth/outlook/callback"]
    state = query["state"][0]

    router.add("POST", "login.microsoftonline.com", FakeResponse(
        {"access_token": "mtok", "refresh_token": "mref", "expires_in": 3600}))
    router.add("GET", "graph.microsoft.com/v1.0/me", FakeResponse(
        {"mail": "student@school.edu", "displayName": "Student"}))
    r = client.get(f"/oauth/outlook/callback?state={state}&code=abc")
    assert r.status_code == 302 and "onedrive_connected" in r.headers["Location"]
    token = _stored_token(uid, "onedrive")
    assert token["access_token"] == "mtok" and token["expires_at"] > time.time()
    with App.app.app_context():
        assert OutlookIntegration.query.filter_by(user_id=uid).count() == 0


def test_microsoft_callback_with_a_bad_state_is_a_redirect_not_a_json_page(client, configured, uid):
    r = client.get("/oauth/outlook/callback?state=forged&code=abc")
    assert r.status_code == 302 and "outlook_error" in r.headers["Location"]


def test_match_searches_drive_refreshes_the_token_and_persists_it(client, configured, uid, router):
    _connect(uid, "google_drive", {"token": "stale", "refresh_token": "r1", "expires_at": time.time() - 60,
                                   "scopes": [gd.SCOPE_FILE, gd.SCOPE_READONLY]})
    router.add("POST", "oauth2.googleapis.com/token", FakeResponse({"access_token": "fresh", "expires_in": 3600}))
    router.add("GET", "/export", FakeResponse(body=b"Chlorophyll drives photosynthesis in the thylakoid."))
    router.add("GET", "/drive/v3/files", FakeResponse({"files": [
        {"id": "doc_11111", "name": "Photosynthesis notes", "mimeType": gd.GOOGLE_DOC,
         "modifiedTime": "2026-09-20T00:00:00Z", "webViewLink": "https://docs.google.com/d/11111"}]}))
    r = client.get("/api/cloud-docs/match?title=Photosynthesis%20lab&course=Biology")
    assert r.status_code == 200, r.get_json()
    data = r.get_json()
    assert data["connected"] == ["google_drive"]
    assert [d["name"] for d in data["documents"]] == ["Photosynthesis notes"]
    assert "text" not in data["documents"][0]  # the browser gets an excerpt, not the file
    assert "photosynthesis" in data["documents"][0]["excerpt"].lower()
    assert _stored_token(uid, "google_drive")["token"] == "fresh"
    assert all(c["headers"]["Authorization"] == "Bearer fresh"
               for c in router.calls if "googleapis.com/drive" in c["url"])


def test_create_writes_a_study_guide_to_onedrive(client, configured, uid, router):
    _connect(uid, "onedrive", {"access_token": "at", "refresh_token": "rt", "expires_at": time.time() + 3600})
    router.add("GET", "/search(q=", FakeResponse({"value": [
        {"id": "ITEM1", "name": "Hamlet notes.txt", "file": {"mimeType": "text/plain"}, "size": 30,
         "lastModifiedDateTime": "2026-09-10T00:00:00Z", "webUrl": "https://onedrive.live.com/n"}]}))
    router.add("GET", "/items/ITEM1/content", FakeResponse(body=b"Hamlet delays revenge; the ghost appears."))
    router.add("PUT", "/me/drive/root:/IntelliPlan/", FakeResponse(
        {"id": "new1", "name": "Hamlet essay study guide.docx", "webUrl": "https://onedrive.live.com/g"}))
    r = client.post("/api/cloud-docs/create", json={
        "provider": "onedrive", "title": "Hamlet essay", "course": "English", "due_date": "2026-10-20"})
    assert r.status_code == 200, r.get_json()
    body = r.get_json()
    assert body["document"]["url"] == "https://onedrive.live.com/g"
    assert body["used_documents"] == 1
    uploaded = router.find("PUT", "/IntelliPlan/")[0]["data"]
    from docx import Document
    text = "\n".join(p.text for p in Document(io.BytesIO(uploaded)).paragraphs)
    assert "Study guide: Hamlet essay" in text and "Hamlet notes.txt" in text


def test_create_asks_for_a_reconnect_when_the_grant_is_gone(client, configured, uid, router):
    _connect(uid, "google_drive", {"token": "t", "refresh_token": "dead", "expires_at": 1})
    router.add("POST", "oauth2.googleapis.com/token", FakeResponse({"error": "invalid_grant"}, status=400))
    r = client.post("/api/cloud-docs/create", json={"title": "Essay", "provider": "google_drive"})
    assert r.status_code == 409 and r.get_json()["connect_required"]


def test_disconnect_removes_only_that_provider(client, configured, uid):
    _connect(uid, "google_drive", {"token": "t"})
    _connect(uid, "onedrive", {"access_token": "t"})
    assert client.post("/api/cloud-docs/disconnect/google_drive").get_json()["status"] == "ok"
    assert _stored_token(uid, "google_drive") is None
    assert _stored_token(uid, "onedrive") is not None


def test_catalog_marks_connected_cloud_providers(client, configured, uid):
    _connect(uid, "onedrive", {"access_token": "t"})
    items = {i["id"]: i for i in client.get("/api/integrations/status").get_json()["integrations"]}
    assert items["onedrive"]["connected"] and items["onedrive"]["detail"] == "s@example.com"
    assert items["google_drive"]["methods"][0]["start_url"] == "/oauth/google-drive"


def test_account_deletion_covers_cloud_docs(client, configured, uid):
    _connect(uid, "google_drive", {"token": "t"})
    r = client.post("/account/delete", json={"confirm": "DELETE"})
    assert r.status_code in (200, 302), r.get_data(as_text=True)[:300]
    assert _stored_token(uid, "google_drive") is None


# Outlook, end to end through the app.

def _connect_outlook(user_id, token):
    with App.app.app_context():
        db.session.add(OutlookIntegration(user_id=user_id, token_data=json.dumps(token),
                                          account_email="s@outlook.com"))
        db.session.commit()


def test_outlook_token_without_expiry_is_refreshed_and_saved(client, configured, uid, router):
    _connect_outlook(uid, {"access_token": "old", "refresh_token": "r"})
    router.add("POST", "login.microsoftonline.com", FakeResponse(
        {"access_token": "new", "expires_in": 3600}))
    router.add("GET", "/me/calendarView", FakeResponse({"value": [
        {"id": "e1", "subject": "Practice", "start": {"dateTime": "2026-10-01T15:00:00"},
         "end": {"dateTime": "2026-10-01T16:00:00"}}]}))
    data = client.get("/calendar/events").get_json()
    assert data["connected"] is True
    assert [e["source"] for e in data["events"]] == ["outlook_calendar"]
    with App.app.app_context():
        stored = json.loads(OutlookIntegration.query.filter_by(user_id=uid).first().token_data)
    assert stored["access_token"] == "new" and stored["refresh_token"] == "r"


def test_calendar_export_pushes_to_outlook_when_only_outlook_is_connected(client, configured, uid, monkeypatch):
    _connect_outlook(uid, {"access_token": "a", "refresh_token": "r", "expires_at": time.time() + 3600})
    monkeypatch.setattr(outlook, "graph_get", lambda *a, **k: {"value": []})
    monkeypatch.setattr(outlook, "graph_post", lambda token, path, body: {"id": "ev"})
    r = client.post("/calendar/export", json={"schedule_data": {"schedule": [{"date": "2026-10-02", "blocks": [
        {"assignment": "Essay", "time_slot": "4:30 PM - 5:30 PM", "duration_minutes": 60}]}]}})
    body = r.get_json()
    assert body["status"] == "ok" and body["created"] == 1
    assert body["providers"]["outlook"] == {"created": 1, "skipped": 0}


def test_calendar_export_with_nothing_connected_says_so(client, configured, uid):
    r = client.post("/calendar/export", json={"schedule_data": {"schedule": []}})
    assert r.get_json() == {"status": "error", "message": "No calendar connected"}


# The tutor, grounded in the student's own documents.

def test_tutor_studies_a_non_canvas_assignment_from_drive(client, configured, uid, monkeypatch):
    from types import SimpleNamespace
    import chatbot_api

    source = dm.Source("google_drive",
                       lambda keywords, course: [{"id": "1", "name": "Cell biology notes",
                                                  "modified": "2026-09-01T00:00:00Z"}],
                       lambda item: "Mitochondria make ATP for the cell. UNIQUE-DRIVE-TEXT")
    monkeypatch.setattr(App, "cloud_docs_sources", lambda user_id=None: [source])
    monkeypatch.setattr(chatbot_api, "_check_and_increment_tutor_limit", lambda: (True, 49, 50))
    monkeypatch.setattr(chatbot_api.ai_firewall, "guard",
                        lambda *a, **k: SimpleNamespace(max_output_tokens=1800, plan="free"))
    monkeypatch.setattr(chatbot_api.ai_firewall, "record_tokens", lambda *a, **k: None)
    monkeypatch.setattr(chatbot_api, "_safety_check_user_message", lambda *a: None)
    monkeypatch.setattr(chatbot_api, "_safety_check_assistant_reply", lambda *a: None)
    monkeypatch.setattr(chatbot_api, "_load_tutor_memory", lambda: {"id": 1, "profile_json": "{}"})
    monkeypatch.setattr(chatbot_api, "_build_personalization_prompt", lambda **k: None)
    monkeypatch.setattr(chatbot_api, "_prepare_adaptive_turn", lambda *a: None)
    monkeypatch.setattr(chatbot_api, "_load_user_identity", lambda: {})
    monkeypatch.setattr(chatbot_api, "_save_tutor_profile", lambda *a: None)
    monkeypatch.setattr(chatbot_api, "_record_adaptive_turn", lambda *a: None)
    monkeypatch.setattr(chatbot_api, "_ensure_conversation", lambda *a: {"id": 5, "title": "Cells"})
    monkeypatch.setattr(chatbot_api, "_save_conversation", lambda *a: None)
    seen = []
    monkeypatch.setattr(chatbot_api, "_llm_chat",
                        lambda **kw: seen.extend(kw["messages"]) or "What does ATP do?")
    r = client.post("/api/tutor", json={
        "messages": [{"role": "user", "content": "Help me start."}],
        "assignment_ref": {"title": "Cell biology quiz", "course": "Biology"},
    })
    assert r.status_code == 200, r.get_json()
    assert r.get_json()["assignment_context"]["files_read"] == ["Your Google Drive: Cell biology notes"]
    assert any("UNIQUE-DRIVE-TEXT" in m["content"] for m in seen)


def test_tutor_non_canvas_ref_without_any_drive_is_a_409(client, unconfigured, uid):
    r = client.post("/api/tutor/assignment-map",
                    json={"assignment_ref": {"title": "Cell biology quiz"}})
    assert r.status_code == 409
