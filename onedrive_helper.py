"""Microsoft identity and OneDrive file operations using delegated Graph access.

Uses the Outlook app registration (``MICROSOFT_CLIENT_ID`` and friends) and
its redirect URI, read through ``outlook_calendar_helper`` so both agree on
the defaults. The token is a separate grant from Outlook Calendar with its
own scopes, so disconnecting one never breaks the other.

Two layers: the selected-file import/edit functions behind /study-files
(token-dict based), and below them the assignment-matching and study-guide
functions that take a self-refreshing :class:`OneDriveClient`.

Why ``Files.ReadWrite`` rather than ``Files.ReadWrite.AppFolder``: the app
folder scope would hide study guides in ``Apps/IntelliPlan``, a folder
students do not know exists.
"""

from __future__ import annotations

import base64
import hashlib
import os
import re
import secrets
from io import BytesIO
from urllib.parse import quote, urlencode

import requests

import outlook_calendar_helper as outlook
from cloud_token_client import ProviderError, TokenClient, TokenExpired, read_capped
from intelliplan.integrations.document_content import CloudDocumentError

AUTHORITY = outlook.AUTHORITY
GRAPH = "https://graph.microsoft.com/v1.0"
SCOPES = ("openid", "profile", "offline_access", "User.Read", "Files.ReadWrite")


def configured() -> bool:
    # Same test as Outlook: the redirect URI has a sensible default now, so
    # requiring it as well kept a correctly registered app hidden.
    return outlook.configured()


def get_auth_url(state: str) -> tuple[str, str]:
    code_verifier = secrets.token_urlsafe(64)
    code_challenge = base64.urlsafe_b64encode(
        hashlib.sha256(code_verifier.encode()).digest()
    ).rstrip(b"=").decode()
    params = {
        "client_id": os.environ["MICROSOFT_CLIENT_ID"],
        "response_type": "code",
        "redirect_uri": outlook.redirect_uri(),
        "response_mode": "query",
        "scope": " ".join(SCOPES),
        "state": state,
        "prompt": "select_account",
        "code_challenge": code_challenge,
        "code_challenge_method": "S256",
    }
    return f"{outlook._authority()}/authorize?{urlencode(params)}", code_verifier


def _token_request(payload: dict) -> dict:
    response = requests.post(f"{outlook._authority()}/token", data={
        "client_id": os.environ["MICROSOFT_CLIENT_ID"],
        "client_secret": os.environ["MICROSOFT_CLIENT_SECRET"],
        "scope": " ".join(SCOPES),
        **payload,
    }, timeout=15)
    response.raise_for_status()
    return response.json()


def exchange_code(code: str, code_verifier: str | None = None) -> dict:
    payload = {
        "grant_type": "authorization_code",
        "code": code,
        "redirect_uri": outlook.redirect_uri(),
    }
    if code_verifier:
        payload["code_verifier"] = code_verifier
    return _token_request(payload)


def refresh_token(token: dict) -> dict:
    refreshed = _token_request({
        "grant_type": "refresh_token",
        "refresh_token": token.get("refresh_token", ""),
    })
    return {**token, **refreshed,
            "refresh_token": refreshed.get("refresh_token") or token.get("refresh_token", "")}


def _headers(token: dict, extra: dict | None = None) -> dict:
    access_token = token.get("access_token")
    if not access_token:
        raise CloudDocumentError("OneDrive is not connected. Please reconnect it.")
    return {"Authorization": f"Bearer {access_token}", **(extra or {})}


def graph_get(token: dict, path: str, params: dict | None = None) -> dict:
    response = requests.get(f"{GRAPH}{path}", params=params, headers=_headers(token), timeout=20)
    response.raise_for_status()
    return response.json()


def profile(token: dict) -> dict:
    return graph_get(token, "/me", {"$select": "displayName,mail,userPrincipalName"})


def search_files(token: dict, query: str = "") -> list[dict]:
    if query.strip():
        escaped = query.strip().replace("'", "''")[:100]
        path = f"/me/drive/root/search(q='{quote(escaped, safe='')}')"
    else:
        path = "/me/drive/root/children"
    result = graph_get(token, path, {"$select": "id,name,file,folder,webUrl,eTag,size,lastModifiedDateTime", "$top": "100"})
    files = result.get("value", [])
    next_url = result.get("@odata.nextLink")
    while next_url and len(files) < 300:
        response = requests.get(next_url, headers=_headers(token), timeout=20)
        response.raise_for_status()
        result = response.json()
        files.extend(result.get("value", []))
        next_url = result.get("@odata.nextLink")
    return files[:300]


def get_file(token: dict, item_id: str) -> dict:
    response = graph_get(token, f"/me/drive/items/{quote(item_id, safe='')}", {
        "$select": "id,name,file,folder,webUrl,eTag,size,lastModifiedDateTime",
    })
    if response.get("folder"):
        raise CloudDocumentError("Choose a file, not a folder.")
    return response


def download_file(token: dict, item_id: str, etag: str | None = None) -> tuple[dict, bytes]:
    info = get_file(token, item_id)
    if int(info.get("size") or 0) > 2 * 1024 * 1024:
        raise CloudDocumentError("This file is larger than IntelliPlan can import (2 MB).")
    response = requests.get(
        f"{GRAPH}/me/drive/items/{quote(item_id, safe='')}/content",
        headers=_headers(token, {"If-Match": etag} if etag else None),
        timeout=30, stream=True,
    )
    if response.status_code == 412:
        response.close()
        raise RuntimeError("This OneDrive file changed. Reload it before saving to avoid overwriting newer edits.")
    try:
        response.raise_for_status()
        from intelliplan.integrations.document_content import read_limited_response
        content = read_limited_response(response)
    except Exception:
        response.close()
        raise
    return info, content


def replace_text_file(token: dict, item_id: str, content: str, etag: str) -> dict:
    info = get_file(token, item_id)
    suffix = (info.get("name") or "").lower().rsplit(".", 1)[-1]
    if suffix not in {"txt", "md", "csv"}:
        raise CloudDocumentError("OneDrive editing is enabled for text, Markdown, and CSV files. Word documents can be imported as study context.")
    response = requests.put(
        f"{GRAPH}/me/drive/items/{quote(item_id, safe='')}/content",
        data=content.encode("utf-8"),
        headers=_headers(token, {"Content-Type": "text/plain; charset=utf-8", "If-Match": etag}),
        timeout=30,
    )
    if response.status_code == 412:
        response.close()
        raise RuntimeError("This OneDrive file changed. Reload it before saving to avoid overwriting newer edits.")
    response.raise_for_status()
    return response.json()


# ── Assignment matching and study guides ─────────────────────────────

FOLDER_NAME = "IntelliPlan"
MAX_DOWNLOAD_BYTES = 2 * 1024 * 1024
MAX_TEXT_CHARS = 12000


class OneDriveClient(TokenClient):
    access_key = "access_token"

    def _refresh(self, token):
        try:
            fresh = _token_request(
                {"grant_type": "refresh_token", "refresh_token": token.get("refresh_token", "")})
        except requests.HTTPError as exc:
            status = getattr(exc.response, "status_code", 0)
            raise TokenExpired(f"Microsoft refresh rejected ({status})") from None
        return fresh


# ── Search ────────────────────────────────────────────────────────────


def search_path(text: str) -> str:
    """``/me/drive/root/search(q='…')`` with the term safely embedded.

    OData string literals escape a single quote by doubling it, and the whole
    thing is then percent-encoded because it sits in the URL path. Getting
    either wrong turns "Romeo & Juliet's balcony" into a 400.
    """
    term = re.sub(r"\s+", " ", str(text or "")).strip()[:120]
    if not term:
        raise ValueError("a search term is required")
    literal = term.replace("'", "''")
    return f"/me/drive/root/search(q='{quote(literal, safe='')}')"


def search_documents(client: OneDriveClient, text: str, limit: int = 15) -> list:
    """Files (not folders) matching ``text`` in name or content, normalised
    for the assignment matcher. :func:`search_files` above is the raw listing
    the /study-files browser pages through."""
    data = client.json("GET", GRAPH + search_path(text), params={
        "$top": str(max(1, min(50, int(limit or 15)))),
        "$select": "id,name,file,folder,size,lastModifiedDateTime,webUrl,parentReference",
    })
    out = []
    for item in data.get("value") or []:
        if not isinstance(item, dict) or not item.get("id") or item.get("folder"):
            continue
        out.append({
            "provider": "onedrive",
            "id": str(item["id"]),
            "name": str(item.get("name") or "Untitled")[:200],
            "mime": str((item.get("file") or {}).get("mimeType") or ""),
            "modified": str(item.get("lastModifiedDateTime") or ""),
            "url": str(item.get("webUrl") or ""),
            "size": int(item.get("size") or 0),
            # Needed to fetch content for items shared from someone else's
            # drive; /me/drive/items/{id} would 404 on those.
            "drive_id": str((item.get("parentReference") or {}).get("driveId") or ""),
        })
    return out


# ── Reading ───────────────────────────────────────────────────────────


def _item_path(item: dict) -> str:
    item_id = str(item.get("id") or "")
    drive_id = str(item.get("drive_id") or "")
    safe = re.compile(r"[A-Za-z0-9!_\-\.]{3,200}")
    if not safe.fullmatch(item_id) or (drive_id and not safe.fullmatch(drive_id)):
        raise ProviderError("invalid OneDrive item id")
    if drive_id:
        return f"{GRAPH}/drives/{drive_id}/items/{item_id}"
    return f"{GRAPH}/me/drive/items/{item_id}"


def fetch_text(client: OneDriveClient, item: dict, max_chars: int = MAX_TEXT_CHARS) -> str:
    """Plain text of one OneDrive file, or ``""`` when it is not cheap to read.

    ``/content`` answers with a 302 to a short-lived pre-authenticated
    download URL on another host. ``requests`` follows it and, because the
    host changes, drops our Authorization header on the way -- which is
    exactly right: that URL needs no token and should never see one.
    """
    from intelliplan.integrations.document_content import extract_text, is_supported

    if not is_supported(item.get("name"), item.get("mime")):
        return ""
    if item.get("size") and int(item["size"]) > MAX_DOWNLOAD_BYTES:
        return ""
    response = client.request("GET", _item_path(item) + "/content", stream=True)
    with response:
        if response.status_code != 200:
            return ""
        body = read_capped(response, MAX_DOWNLOAD_BYTES)
    if body is None:
        return ""
    try:
        text = extract_text(item.get("name") or "", item.get("mime") or "", body)
    except Exception as exc:
        print(f"[onedrive] could not extract text: {type(exc).__name__}")
        return ""
    return re.sub(r"\s+", " ", text or "").strip()[:max_chars]


# ── Writing ───────────────────────────────────────────────────────────


def _safe_filename(title: str, ext: str) -> str:
    # OneDrive rejects " * : < > ? / \ | and names ending in a dot or space.
    name = re.sub(r'[\\/:*?"<>|#%]+', " ", str(title or "Study guide"))
    name = re.sub(r"\s+", " ", name).strip(" .")[:120] or "Study guide"
    return f"{name}.{ext}"


def markdown_to_docx(markdown: str, title: str = "") -> bytes | None:
    """A real Word document from the study-guide Markdown, via python-docx
    (already a dependency for reading .docx). None if it is unavailable, in
    which case the caller uploads the Markdown as-is."""
    try:
        from docx import Document
    except Exception:
        return None
    doc = Document()
    for raw in str(markdown or "").splitlines():
        line = raw.rstrip()
        heading = re.match(r"^(#{1,3})\s+(.*)$", line)
        bullet = re.match(r"^\s*[-*]\s+(\[[ xX]\]\s+)?(.*)$", line)
        if heading:
            doc.add_heading(_plain(heading.group(2)), level=len(heading.group(1)))
        elif bullet:
            prefix = "☐ " if bullet.group(1) else ""
            doc.add_paragraph(prefix + _plain(bullet.group(2)), style="List Bullet")
        elif line.strip():
            doc.add_paragraph(_plain(line.strip()))
    if title:
        doc.core_properties.title = title[:200]
    out = BytesIO()
    doc.save(out)
    return out.getvalue()


def _plain(text: str) -> str:
    text = re.sub(r"\*\*(.+?)\*\*", r"\1", text)
    return re.sub(r"\[([^\]]+)\]\((https://[^)\s]+)\)", r"\1 (\2)", text)


def upload_path(filename: str) -> str:
    """Path-addressed upload into /IntelliPlan, creating the folder implicitly.

    Graph creates missing parent folders on a path-based PUT, so there is no
    separate "make the folder" round trip. ``conflictBehavior=rename`` makes a
    second guide for the same assignment "Name 1.docx" instead of silently
    overwriting the one the student may have written in.
    """
    return (f"{GRAPH}/me/drive/root:/{quote(FOLDER_NAME)}/{quote(filename)}:/content"
            "?@microsoft.graph.conflictBehavior=rename")


def create_document(client: OneDriveClient, title: str, markdown: str) -> dict:
    """Upload a study guide as .docx (or .md) into /IntelliPlan."""
    body = markdown_to_docx(markdown, title)
    if body is not None:
        filename = _safe_filename(title, "docx")
        content_type = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
    else:
        filename = _safe_filename(title, "md")
        body = str(markdown or "").encode("utf-8")
        content_type = "text/markdown"
    # Simple upload is limited to 250 MB; a study guide is a few kilobytes.
    response = client.request("PUT", upload_path(filename), data=body,
                              headers={"Content-Type": content_type})
    if response.status_code >= 400:
        raise ProviderError(f"OneDrive upload failed with HTTP {response.status_code}",
                            response.status_code)
    data = response.json()
    return {"provider": "onedrive", "id": data.get("id"),
            "name": data.get("name") or filename, "url": data.get("webUrl") or ""}
